"""Bandeja de salida (modo de escritura BANDEJA), con conectores falsos.

Los conectores falsos leen como el CSV pero NUNCA deben escribir en el "ERP":
en modo bandeja confirmar solo encola la operación.
"""

from datetime import timedelta
from pathlib import Path

import pytest
from mcp import Client
from openpyxl import load_workbook

from stockmind.adapters.csv_excel import CsvExcelConnector
from stockmind.application.bandeja import MemoriaBandejaStore, SqliteBandejaStore
from stockmind.application.confirmaciones import GestorConfirmaciones, MemoriaConfirmacionStore
from stockmind.application.servicio import InventarioService
from stockmind.domain.bandeja import Coincidencia, EstadoBandeja, OperacionBandeja
from stockmind.domain.errores import (
    BandejaNoDisponible,
    EscrituraNoHabilitada,
    ErrorValidacion,
    OperacionNoEncontrada,
    TransicionInvalida,
)
from stockmind.domain.modelos import ErrorCampo, ModoEscritura, Revision
from stockmind.mcp_server.server import crear_servidor
from stockmind.ports.bandeja import IntegracionBandeja
from tests.conftest import HOY

PRODUCTO = {"nombre": "Martillo de uña", "categoria": "Herramientas", "unidad": "unidad",
            "precio_venta": 65, "costo": 45}


class ConectorBandejaFalso(CsvExcelConnector):
    """ERP de solo lectura: lee del CSV, no escribe nunca."""

    nombre = "erp_falso"
    nombre_visible = "ERP de prueba"
    modo_escritura = ModoEscritura.BANDEJA

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.revisiones: list[tuple[str, int]] = []
        self.revision_extra = Revision()

    def crear_producto(self, datos):
        raise AssertionError("en modo bandeja no se escribe en el ERP")

    def registrar_movimiento(self, datos):
        raise AssertionError("en modo bandeja no se escribe en el ERP")

    def advertencias_producto_nuevo(self, datos):
        return [f"El ERP guardará el nombre como {datos.nombre.upper()}"]

    def validar_movimiento_nuevo(self, movimiento, producto, pendientes):
        self.revisiones.append((producto.id, len(pendientes)))
        revision = self.revision_extra.model_copy(deep=True)
        revision.detalle.setdefault("almacen", movimiento.almacen or "Principal")
        return revision


class ConectorConIntegracion(ConectorBandejaFalso, IntegracionBandeja):
    """Además trae su propia exportación, guía y conciliación."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.en_erp: set[str] = set()        # nombres que "ya están" en el ERP
        self.posibles: dict[str, str] = {}   # producto_id → descripción de una posible coincidencia

    def guia(self, operacion):
        return [f"Abrir la pantalla del ERP para {operacion.referencia}", "Guardar"]

    def exportar(self, operaciones, tipo, carpeta, momento):
        ruta = Path(carpeta) / f"{tipo}_erp.txt"
        ruta.parent.mkdir(parents=True, exist_ok=True)
        ruta.write_text("\n".join(op.referencia for op in operaciones), encoding="utf-8")
        return [ruta]

    def verificar(self, operacion):
        if operacion.tipo == "crear_producto" and operacion.datos["nombre"].upper() in self.en_erp:
            return [Coincidencia(exacta=True, descripcion="Encontrado en el ERP", id_externo="99")]
        if operacion.tipo == "registrar_movimiento" and operacion.datos["producto_id"] in self.posibles:
            return [Coincidencia(exacta=False, descripcion=self.posibles[operacion.datos["producto_id"]])]
        return []


def _servicio(conector, reloj, auditoria, tmp_path, store=None) -> InventarioService:
    return InventarioService(
        conector, GestorConfirmaciones(MemoriaConfirmacionStore(), ttl=timedelta(minutes=10), reloj=reloj),
        auditoria, bandeja=store or MemoriaBandejaStore(), exports_dir=tmp_path / "exports",
        origen="test", hoy=lambda: HOY, reloj=reloj,
    )


@pytest.fixture
def falso(datos_dir) -> ConectorBandejaFalso:
    return ConectorBandejaFalso(datos_dir)


@pytest.fixture
def integrado(datos_dir) -> ConectorConIntegracion:
    return ConectorConIntegracion(datos_dir)


@pytest.fixture
def servicio_bandeja(falso, reloj, auditoria, tmp_path) -> InventarioService:
    return _servicio(falso, reloj, auditoria, tmp_path)


@pytest.fixture
def servicio_integrado(integrado, reloj, auditoria, tmp_path) -> InventarioService:
    return _servicio(integrado, reloj, auditoria, tmp_path)


def _crear(servicio, datos=PRODUCTO) -> dict:
    return servicio.confirmar(servicio.preparar_crear_producto(datos)["token"])


def _salida(servicio, cantidad, **extra) -> dict:
    datos = {"producto_id": "P0002", "tipo": "salida", "cantidad": cantidad, "motivo": "Venta mostrador"} | extra
    return servicio.preparar_movimiento(datos)


# --------------------------------------------------------- confirmar → encolar

def test_vista_previa_avisa_que_quedara_pendiente(servicio_bandeja):
    vista = servicio_bandeja.preparar_crear_producto(PRODUCTO)
    assert vista["modo_escritura"] == "bandeja"
    assert "PENDIENTE" in vista["instruccion"] and "ERP de prueba" in vista["instruccion"]
    assert "MARTILLO DE UÑA" in vista["advertencias"][0]


def test_confirmar_deja_pendiente_y_no_escribe_en_el_erp(servicio_bandeja, datos_dir, auditoria):
    respuesta = _crear(servicio_bandeja)
    assert respuesta["estado"] == "pendiente_en_erp" and respuesta["referencia"] == "SM-000001"
    assert "NO se registró" in respuesta["mensaje"] and "PENDIENTE" in respuesta["mensaje"]
    assert respuesta["guia"] and "SM-000001" in " ".join(respuesta["guia"])   # guía genérica
    assert CsvExcelConnector(datos_dir).buscar_productos("martillo") == []     # el ERP no cambió

    (op,) = servicio_bandeja.listar_pendientes()["operaciones"]
    assert op["referencia"] == "SM-000001" and op["estado"] == "pendiente"
    assert op["descripcion"] == "Crear producto «Martillo de uña»" and op["guia"]
    eventos = [(r["evento"], r["referencia"]) for r in auditoria.leer()]
    assert eventos == [("preparada", None), ("encolada", "SM-000001")]


def test_movimiento_pendiente_guarda_el_detalle_del_erp(servicio_bandeja, falso):
    falso.revision_extra = Revision(advertencias=["Se usará el lote L-1"], detalle={"lotes": [{"lote": "L-1"}]})
    vista = _salida(servicio_bandeja, 2)
    assert vista["vista_previa"]["detalle_erp"]["lotes"] == [{"lote": "L-1"}]
    assert "Se usará el lote L-1" in vista["advertencias"]
    respuesta = servicio_bandeja.confirmar(vista["token"])
    op = servicio_bandeja.bandeja.obtener(1)
    assert op.detalle_erp["lotes"] == [{"lote": "L-1"}] and op.datos["cantidad"] == 2
    assert "de «Tubería PVC 1/2\"»" in servicio_bandeja.listar_pendientes()["operaciones"][0]["descripcion"]
    assert respuesta["vista_previa"]["stock_resultante"] == 3


def test_reglas_del_erp_bloquean_antes_del_token(servicio_bandeja, falso, auditoria):
    falso.revision_extra = Revision(errores=[ErrorCampo(campo="lote", mensaje="no existe el lote X")])
    with pytest.raises(ErrorValidacion, match="no existe el lote X"):
        _salida(servicio_bandeja, 1)
    assert auditoria.leer() == []


def test_stock_comprometido_por_salidas_pendientes(servicio_bandeja, falso):
    servicio_bandeja.confirmar(_salida(servicio_bandeja, 3)["token"])   # stock 5, quedan 2 libres
    vista = _salida(servicio_bandeja, 2)
    assert vista["vista_previa"]["stock_comprometido_en_bandeja"] == 3
    assert vista["vista_previa"]["stock_resultante"] == 0
    assert falso.revisiones[-1] == ("P0002", 1)                         # el conector ve la pendiente
    with pytest.raises(ErrorValidacion, match="3 ya comprometidas"):
        _salida(servicio_bandeja, 3)


def test_salida_descartada_libera_el_stock(servicio_bandeja):
    servicio_bandeja.confirmar(_salida(servicio_bandeja, 5)["token"])
    with pytest.raises(ErrorValidacion):
        _salida(servicio_bandeja, 1)
    servicio_bandeja.descartar_pendiente("SM-000001", "Se anuló la venta")
    assert _salida(servicio_bandeja, 1)["ok"]


def test_no_duplica_productos_pendientes(servicio_bandeja):
    primera = servicio_bandeja.preparar_crear_producto(PRODUCTO)
    segunda = servicio_bandeja.preparar_crear_producto(PRODUCTO | {"nombre": "martillo de UÑA"})
    servicio_bandeja.confirmar(primera["token"])
    with pytest.raises(ErrorValidacion, match="SM-000001"):              # revalidación al confirmar
        servicio_bandeja.confirmar(segunda["token"])
    with pytest.raises(ErrorValidacion, match="pendiente de registrar"):
        servicio_bandeja.preparar_crear_producto(PRODUCTO)


def test_modo_ninguna_rechaza_antes_de_validar(datos_dir, reloj, auditoria, tmp_path):
    conector = ConectorBandejaFalso(datos_dir)
    conector.modo_escritura = ModoEscritura.NINGUNA
    servicio = _servicio(conector, reloj, auditoria, tmp_path)
    with pytest.raises(EscrituraNoHabilitada):
        servicio.preparar_crear_producto({})
    with pytest.raises(BandejaNoDisponible, match="no acepta escrituras"):
        servicio.exportar_pendientes()


def test_modo_directa_no_usa_bandeja(conector, reloj, auditoria, tmp_path):
    servicio = _servicio(conector, reloj, auditoria, tmp_path)        # CSV: escribe directo
    assert servicio.listar_pendientes()["total"] == 0
    with pytest.raises(BandejaNoDisponible, match="escribe directamente"):
        servicio.verificar_pendientes()
    vista = servicio.preparar_movimiento({"producto_id": "P0003", "tipo": "entrada", "cantidad": 2})
    assert "PENDIENTE" not in vista["instruccion"]
    assert servicio.confirmar(vista["token"])["resultado"]["saldo"] == 2   # sí escribió
    assert servicio.listar_pendientes()["total"] == 0


# ------------------------------------------------------------------ estados

def test_marcar_registrada_una_sola_vez(servicio_bandeja, auditoria, reloj):
    _crear(servicio_bandeja)
    reloj.avanzar(hours=2)
    hecha = servicio_bandeja.marcar_registrada("SM-000001", "Importado en el ERP")
    assert hecha["operacion"]["estado"] == "registrada"
    assert hecha["operacion"]["registrada_por"] == "manual"
    assert hecha["operacion"]["observacion"] == "Importado en el ERP"
    assert "guia" not in hecha["operacion"]
    with pytest.raises(TransicionInvalida, match="ya está registrada"):
        servicio_bandeja.marcar_registrada(1)
    with pytest.raises(TransicionInvalida):
        servicio_bandeja.descartar_pendiente("sm-1", "tarde")
    ultimo = auditoria.leer()[-1]
    assert ultimo["evento"] == "registrada" and ultimo["referencia"] == "SM-000001"
    assert ultimo["datos"] == {"de": "pendiente", "a": "registrada", "observacion": "Importado en el ERP",
                               "registrada_por": "manual"}


def test_descartar_exige_motivo(servicio_bandeja, auditoria):
    _crear(servicio_bandeja)
    with pytest.raises(ErrorValidacion, match="motivo"):
        servicio_bandeja.descartar_pendiente("SM-000001", "  ")
    hecha = servicio_bandeja.descartar_pendiente("SM-000001", "El cliente desistió")
    assert hecha["operacion"]["estado"] == "descartada"
    assert auditoria.leer()[-1]["evento"] == "descartada"
    assert servicio_bandeja.listar_pendientes()["total"] == 0
    assert servicio_bandeja.listar_pendientes("descartada")["total"] == 1
    assert servicio_bandeja.listar_pendientes("todas", "productos")["total"] == 1


@pytest.mark.parametrize("id_malo, error", [("abc", ErrorValidacion), ("SM-000009", OperacionNoEncontrada)])
def test_ids_invalidos(servicio_bandeja, id_malo, error):
    with pytest.raises(error):
        servicio_bandeja.marcar_registrada(id_malo)


def test_filtros_invalidos(servicio_bandeja):
    with pytest.raises(ErrorValidacion, match="pendiente, registrada"):
        servicio_bandeja.listar_pendientes("aprobada")
    with pytest.raises(ErrorValidacion, match="productos, movimientos"):
        servicio_bandeja.exportar_pendientes("clientes")


# ---------------------------------------------------------------- exportar

def test_exportacion_generica(servicio_bandeja, tmp_path, auditoria):
    assert servicio_bandeja.exportar_pendientes()["archivos"] == []
    _crear(servicio_bandeja)
    servicio_bandeja.confirmar(_salida(servicio_bandeja, 1)["token"])
    respuesta = servicio_bandeja.exportar_pendientes()
    assert respuesta["operaciones"] == ["SM-000001", "SM-000002"] and len(respuesta["archivos"]) == 2
    assert "Siguen PENDIENTES" in respuesta["mensaje"]

    libro = load_workbook(next(a for a in respuesta["archivos"] if "crear_producto" in a))
    assert libro.sheetnames == ["Pendientes", "Guía"]
    filas = list(libro["Pendientes"].values)
    assert filas[0][:3] == ("referencia", "creada_en", "nombre") and filas[1][0] == "SM-000001"

    for op in servicio_bandeja.listar_pendientes()["operaciones"]:
        assert op["estado"] == "pendiente" and op["exportada_en"] and op["archivo"]
    assert [r["evento"] for r in auditoria.leer()].count("exportada") == 2


def test_exportar_solo_un_tipo(servicio_integrado):
    _crear(servicio_integrado)
    servicio_integrado.confirmar(_salida(servicio_integrado, 1)["token"])
    respuesta = servicio_integrado.exportar_pendientes("movimientos")
    assert respuesta["operaciones"] == ["SM-000002"]
    assert Path(respuesta["archivos"][0]).read_text(encoding="utf-8") == "SM-000002"


def test_solo_se_exportan_las_del_conector_activo(servicio_bandeja):
    servicio_bandeja.bandeja.agregar(OperacionBandeja(
        id=0, tipo="crear_producto", conector="otro_erp", origen="test", datos={"nombre": "X"},
        creada_en=servicio_bandeja.reloj(), actualizada_en=servicio_bandeja.reloj()))
    assert servicio_bandeja.exportar_pendientes()["operaciones"] == []
    assert servicio_bandeja.listar_pendientes()["total"] == 1      # listar muestra todas


# --------------------------------------------------------------- verificar

def test_verificar_marca_exactas_e_informa_posibles(servicio_integrado, integrado, auditoria):
    _crear(servicio_integrado)
    _crear(servicio_integrado, PRODUCTO | {"nombre": "Serrucho"})
    servicio_integrado.confirmar(_salida(servicio_integrado, 1)["token"])
    integrado.en_erp = {"MARTILLO DE UÑA"}
    integrado.posibles = {"P0002": "Ajuste AJU-000004 del mismo día y cantidad"}

    resultado = servicio_integrado.verificar_pendientes()
    assert [r["referencia"] for r in resultado["registradas"]] == ["SM-000001"]
    assert resultado["registradas"][0]["id_externo"] == "99"
    assert resultado["posibles"] == [{"referencia": "SM-000003",
                                      "coincidencias": [{"descripcion": "Ajuste AJU-000004 del mismo día y cantidad"}]}]
    assert resultado["siguen_pendientes"] == ["SM-000002", "SM-000003"]
    assert "NO se marcaron solas" in resultado["mensaje"]

    op = servicio_integrado.bandeja.obtener(1)
    assert op.estado == EstadoBandeja.REGISTRADA and op.registrada_por == "conciliacion" and op.id_externo == "99"
    assert servicio_integrado.bandeja.obtener(3).estado == EstadoBandeja.PENDIENTE   # posible ≠ registrada
    assert auditoria.leer()[-1]["datos"]["registrada_por"] == "conciliacion"


def test_guia_propia_del_conector(servicio_integrado):
    assert _crear(servicio_integrado)["guia"] == ["Abrir la pantalla del ERP para SM-000001", "Guardar"]


def test_verificar_sin_integracion_no_marca_nada(servicio_bandeja):
    _crear(servicio_bandeja)
    resultado = servicio_bandeja.verificar_pendientes()
    assert resultado["registradas"] == [] and resultado["siguen_pendientes"] == ["SM-000001"]


# -------------------------------------------------------------- almacenes

@pytest.fixture(params=["memoria", "sqlite"])
def store(request, tmp_path):
    return MemoriaBandejaStore() if request.param == "memoria" else SqliteBandejaStore(tmp_path / "b d" / "sm.db")


def _op(reloj, **cambios) -> OperacionBandeja:
    return OperacionBandeja(id=0, tipo="crear_producto", conector="erp_falso", origen="test",
                            datos={"nombre": "X", "precio": 1.5}, vista_previa={"nombre": "X"},
                            detalle_erp={"lotes": [{"lote": "A", "cantidad": 2}]},
                            creada_en=reloj(), actualizada_en=reloj()).model_copy(update=cambios)


def test_store_ciclo_completo(store, reloj):
    a, b = store.agregar(_op(reloj)), store.agregar(_op(reloj, tipo="registrar_movimiento"))
    assert (a.id, b.id, a.referencia) == (1, 2, "SM-000001")
    store.actualizar_guia(1, ["paso 1"])
    assert store.obtener(1).guia == ["paso 1"] and store.obtener(1).detalle_erp["lotes"][0]["lote"] == "A"
    assert [o.id for o in store.listar(EstadoBandeja.PENDIENTE, "registrar_movimiento")] == [2]
    reloj.avanzar(minutes=5)
    hecha = store.transicionar(1, EstadoBandeja.REGISTRADA, reloj(), "ok", "manual", "123")
    assert hecha.estado == EstadoBandeja.REGISTRADA and hecha.id_externo == "123"
    assert hecha.actualizada_en > hecha.creada_en
    assert store.transicionar(1, EstadoBandeja.DESCARTADA, reloj(), "x") is None   # ya no está pendiente
    assert store.transicionar(77, EstadoBandeja.DESCARTADA, reloj(), "x") is None
    store.marcar_exportadas([2], reloj(), "exports/a.xlsx")
    assert store.obtener(2).archivo == "exports/a.xlsx" and store.obtener(2).exportada_en is not None
    assert store.obtener(99) is None


def test_sqlite_compartido_entre_procesos(tmp_path, reloj):
    ruta = tmp_path / "stockmind.db"
    mcp, web = SqliteBandejaStore(ruta), SqliteBandejaStore(ruta)
    op = mcp.agregar(_op(reloj))
    assert web.obtener(op.id).datos == {"nombre": "X", "precio": 1.5}
    assert web.transicionar(op.id, EstadoBandeja.REGISTRADA, reloj(), "web") is not None
    assert mcp.transicionar(op.id, EstadoBandeja.DESCARTADA, reloj(), "mcp") is None


# --------------------------------------------------------------------- MCP

@pytest.fixture
def anyio_backend() -> str:
    return "asyncio"


async def _llamar(client, tool, argumentos=None) -> dict:
    resultado = await client.call_tool(tool, argumentos or {})
    assert not resultado.is_error, resultado.content
    return resultado.structured_content


@pytest.mark.anyio
async def test_flujo_completo_por_mcp(servicio_integrado, integrado):
    async with Client(crear_servidor(servicio_integrado)) as client:
        vista = await _llamar(client, "crear_producto", {"datos": PRODUCTO})
        assert vista["modo_escritura"] == "bandeja"
        confirmado = await _llamar(client, "confirmar_operacion", {"token": vista["token"]})
        assert confirmado["estado"] == "pendiente_en_erp" and "NO se registró" in confirmado["mensaje"]

        pendientes = await _llamar(client, "listar_pendientes")
        assert pendientes["operaciones"][0]["referencia"] == "SM-000001"

        exportado = await _llamar(client, "exportar_pendientes", {"tipo": "productos"})
        assert exportado["operaciones"] == ["SM-000001"]

        nada = await _llamar(client, "verificar_pendientes")
        assert nada["siguen_pendientes"] == ["SM-000001"]
        integrado.en_erp = {"MARTILLO DE UÑA"}
        verificado = await _llamar(client, "verificar_pendientes")
        assert verificado["registradas"][0]["referencia"] == "SM-000001"

        otra = await _llamar(client, "marcar_registrada", {"id": "SM-000001"})
        assert otra["ok"] is False and otra["error"] == "TransicionInvalida"
        sin_motivo = await _llamar(client, "descartar_pendiente", {"id": "SM-000001"})
        assert sin_motivo["ok"] is False and sin_motivo["faltantes"][0]["campo"] == "motivo"
