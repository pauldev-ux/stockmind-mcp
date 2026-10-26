"""Reportes con datos controlados: cada número se puede calcular a mano.

Hoy = 25/10/2026. Ventana de 30 días = 26/09 a 25/10.

| Producto  | Categoría    | costo | precio | stock | mín | vendido (30 d)       | otros movimientos (30 d)      |
|-----------|--------------|-------|--------|-------|-----|----------------------|-------------------------------|
| Alfa      | Herramientas | 10    | 20     | 20    | 30  | 30 el 01/10 + 30 el 20/10 = 60 (1200) | entrada 50 el 10/10 |
| Beta      | Herramientas | 5     | 10     | 100   | 10  | 72 el 01/10 (720)    | merma 8 el 02/10              |
| Gamma     | Pinturas     | 2     | 3      | 50    | 0   | 120 el 05/10 (360)   | —                             |
| Delta     | Pinturas     | 1     | 2      | 8     | 10  | 60 el 12/10 (120, almacén Sucursal) | —              |
| Épsilon   | Plomería     | 4     | 6      | 40    | 5   | —  (vendió 5 el 01/05/2026)            | —            |
| Zeta      | Plomería     | 3     | 5      | 10    | 50  | inactivo: no entra en los reportes      |              |

ABC: total 2400 → Alfa 50 %, Beta 30 % (acum. 80), Gamma 15 % (acum. 95), Delta 5 % (acum. 100).
"""

from datetime import date, timedelta
from pathlib import Path

import pandas as pd
import pytest
from mcp import Client
from openpyxl import load_workbook

from stockmind.adapters.csv_excel import CsvExcelConnector, escribir_tabla, ruta_tabla
from stockmind.application.exportar_reportes import exportar_reporte
from stockmind.application.reportes import ServicioReportes, numero
from stockmind.domain.errores import ErrorValidacion
from stockmind.mcp_server.server import crear_servidor
from stockmind.ports.inventario import InventarioConnector
from tests.conftest import HOY

PRODUCTOS = [
    # id, codigo, nombre, categoria, unidad, precio, costo, stock, min, activo
    ("P0001", "ALF", "Alfa", "Herramientas", "unidad", "20", "10", "20", "30", "1"),
    ("P0002", "BET", "Beta", "Herramientas", "unidad", "10", "5", "100", "10", "1"),
    ("P0003", "GAM", "Gamma", "Pinturas", "litro", "3", "2", "50", "0", "1"),
    ("P0004", "DEL", "Delta", "Pinturas", "litro", "2", "1", "8", "10", "1"),
    ("P0005", "EPS", "Épsilon", "Plomería", "pieza", "6", "4", "40", "5", "1"),
    ("P0006", "ZET", "Zeta", "Plomería", "pieza", "5", "3", "10", "50", "0"),
]
# (venta, fecha, almacén, [(producto, cantidad, precio)])
VENTAS = [
    ("V1", "2026-10-01", "Principal", [("P0001", 30, 20), ("P0002", 72, 10)]),
    ("V2", "2026-10-05", "Principal", [("P0003", 120, 3)]),
    ("V3", "2026-10-12", "Sucursal", [("P0004", 60, 2)]),
    ("V4", "2026-10-20", "Principal", [("P0001", 30, 20)]),
    ("V5", "2026-05-01", "Principal", [("P0005", 5, 6)]),
]
# (producto, tipo, sentido, cantidad, fecha); además, una salida por cada línea de venta.
OTROS_MOVIMIENTOS = [
    ("P0001", "entrada", "+", 50, "2026-10-10"),
    ("P0002", "ajuste", "-", 8, "2026-10-02"),
]


def escribir_datos(carpeta: Path) -> None:
    productos = pd.DataFrame(PRODUCTOS, columns=["id", "codigo", "nombre", "categoria", "unidad", "precio_venta",
                                                 "costo", "stock_actual", "stock_minimo", "activo"]).assign(almacen="Principal")
    ventas, detalle, movimientos = [], [], []
    for venta, fecha_venta, almacen, lineas in VENTAS:
        ventas.append((venta, fecha_venta, almacen, str(sum(q * p for _, q, p in lineas))))
        for pid, q, p in lineas:
            detalle.append((venta, pid, str(q), str(p), str(q * p)))
            movimientos.append((pid, "salida", "-", q, fecha_venta, f"Venta {venta}"))
    movimientos += [(pid, t, s, q, f, "otro") for pid, t, s, q, f in OTROS_MOVIMIENTOS]
    movimientos = pd.DataFrame(
        [(f"M{i:06d}", pid, t, s, str(q), f, m) for i, (pid, t, s, q, f, m) in enumerate(movimientos, 1)],
        columns=["id", "producto_id", "tipo", "sentido", "cantidad", "fecha", "motivo"])
    tablas = {
        "productos": productos, "movimientos": movimientos,
        "ventas": pd.DataFrame(ventas, columns=["id", "fecha", "almacen", "total"]),
        "ventas_detalle": pd.DataFrame(detalle, columns=["venta_id", "producto_id", "cantidad", "precio_unitario",
                                                          "subtotal"]),
    }
    for nombre, df in tablas.items():
        escribir_tabla(df, ruta_tabla(carpeta, nombre, "csv"), nombre)


@pytest.fixture
def conector_reportes(tmp_path) -> CsvExcelConnector:
    escribir_datos(tmp_path / "datos")
    return CsvExcelConnector(tmp_path / "datos")


@pytest.fixture
def reportes(conector_reportes) -> ServicioReportes:
    return ServicioReportes(conector_reportes, hoy=lambda: HOY)


def _por_id(filas) -> dict:
    return {f.producto_id: f for f in filas}


# ------------------------------------------------------------------ período

def test_periodos(reportes):
    p = reportes.periodo(None, None, 30, 90)
    assert (p.desde, p.hasta, p.dias) == (date(2026, 9, 26), HOY, 30)
    assert reportes.periodo(None, None, None, 90).dias == 90
    assert reportes.periodo(date(2026, 10, 1), date(2026, 10, 10), None, 30).dias == 10
    with pytest.raises(ErrorValidacion, match="no los dos"):
        reportes.periodo(date(2026, 10, 1), None, 30, 30)
    with pytest.raises(ErrorValidacion, match="futura"):
        reportes.periodo(None, HOY + timedelta(days=1), None, 30)
    with pytest.raises(ErrorValidacion, match="posterior"):
        reportes.periodo(HOY, date(2026, 10, 1), None, 30)


def test_formato_boliviano():
    assert numero(12345.6) == "12.345,6" and numero(1200.0) == "1.200" and numero(None) == "—"


# ------------------------------------------------------------ más vendidos

def test_mas_vendidos(reportes):
    r = reportes.mas_vendidos(ultimos_dias=30)
    assert (r.total_unidades, r.total_monto, r.productos_con_ventas) == (312, 2400, 4)
    assert [f.nombre for f in r.por_cantidad] == ["Gamma", "Beta", "Alfa", "Delta"]   # empate 60: por nombre
    assert [f.nombre for f in r.por_monto] == ["Alfa", "Beta", "Gamma", "Delta"]
    alfa = r.por_monto[0]
    assert (alfa.monto, alfa.participacion_monto, alfa.numero_ventas) == (1200, 50.0, 2)
    assert r.por_cantidad[0].participacion_cantidad == round(120 / 312 * 100, 2)
    assert not r.sin_datos and "26/09/2026" in r.metodo[0]
    assert len(reportes.mas_vendidos(ultimos_dias=30, top=2).por_monto) == 2


def test_mas_vendidos_por_categoria_y_almacen(reportes):
    assert [f.nombre for f in reportes.mas_vendidos(ultimos_dias=30, categoria="pinturas").por_monto] == ["Gamma", "Delta"]
    sucursal = reportes.mas_vendidos(ultimos_dias=30, almacen="sucursal")
    assert [f.nombre for f in sucursal.por_monto] == ["Delta"] and sucursal.total_monto == 120


def test_mas_vendidos_sin_ventas_lo_dice(reportes):
    r = reportes.mas_vendidos(desde=date(2026, 6, 1), hasta=date(2026, 6, 30))
    assert r.sin_datos and r.por_monto == [] and "No hay ventas registradas del 01/06/2026" in r.avisos[0]


def test_categoria_inexistente(reportes):
    r = reportes.mas_vendidos(ultimos_dias=30, categoria="Jardín")
    assert r.sin_datos and "Herramientas" in r.avisos[0]


# -------------------------------------------------------------- bajo stock

def test_bajo_stock_cobertura_y_sugerido(reportes):
    r = reportes.bajo_stock()
    assert [f.nombre for f in r.productos] == ["Delta", "Alfa"]           # menor cobertura primero
    delta, alfa = r.productos
    # Delta: 60 vendidas / 30 = 2 por día; 8 / 2 = 4 días; objetivo máx(10, 2 × 30) = 60 → 60 − 8 = 52
    assert (delta.venta_diaria, delta.cobertura_dias, delta.objetivo, delta.sugerido_reponer) == (2, 4, 60, 52)
    # Alfa: 60 / 30 = 2; 20 / 2 = 10 días; objetivo máx(30, 60) = 60 → 40
    assert (alfa.venta_diaria, alfa.cobertura_dias, alfa.sugerido_reponer, alfa.motivo) == (2, 10, 40, "bajo_minimo")
    assert "Zeta" not in [f.nombre for f in r.productos]                   # inactivo
    assert not r.sin_datos


def test_bajo_stock_con_cobertura_baja(reportes):
    filas = _por_id(reportes.bajo_stock(cobertura_menor_a=50).productos)
    # Gamma: 120 / 30 = 4 por día; 50 / 4 = 12,5 días; objetivo 120 → reponer 70
    assert (filas["P0003"].cobertura_dias, filas["P0003"].sugerido_reponer, filas["P0003"].motivo) == (
        12.5, 70, "cobertura_baja")
    # Beta: 72 / 30 = 2,4; 100 / 2,4 = 41,7 días; objetivo 72 < 100 → no hay que reponer
    assert (filas["P0002"].cobertura_dias, filas["P0002"].sugerido_reponer) == (41.7, 0)


def test_bajo_stock_sin_ventas_no_inventa_ceros(conector_reportes):
    sin_ventas = ServicioReportes(conector_reportes, hoy=lambda: date(2027, 3, 1))
    r = sin_ventas.bajo_stock()
    assert r.sin_datos and "No hay ventas registradas" in r.avisos[0]
    filas = _por_id(r.productos)
    assert filas["P0001"].venta_diaria is None and filas["P0001"].cobertura_dias is None
    assert filas["P0001"].sugerido_reponer == 10 and filas["P0004"].sugerido_reponer == 2   # solo hasta el mínimo


def test_cobertura_indefinida_si_el_producto_no_se_vende(reportes):
    eps = _por_id(reportes.bajo_stock(cobertura_menor_a=10_000).productos)
    assert "P0005" not in eps    # sin ventas: cobertura indefinida, no "baja"


# ---------------------------------------------------------- sin movimiento

def test_sin_movimiento(reportes):
    r = reportes.sin_movimiento()
    assert [f.nombre for f in r.productos] == ["Épsilon"]
    eps = r.productos[0]
    assert (eps.capital_inmovilizado, eps.ultima_salida) == (160, date(2026, 5, 1))
    assert eps.dias_sin_movimiento == (HOY - date(2026, 5, 1)).days == 177
    assert r.capital_por_moneda == {"BOB": 160}


def test_sin_movimiento_ventana_larga(reportes):
    r = reportes.sin_movimiento(dias=200)
    assert r.productos == [] and "tuvieron ventas o salidas en los últimos 200 días" in r.avisos[-1]


# --------------------------------------------------------------------- ABC

def test_abc_calculado_a_mano(reportes):
    r = reportes.abc(ultimos_dias=30)
    filas = [(f.nombre, f.porcentaje, f.porcentaje_acumulado, f.clase) for f in r.productos]
    assert filas == [("Alfa", 50, 50, "A"), ("Beta", 30, 80, "A"), ("Gamma", 15, 95, "B"), ("Delta", 5, 100, "C")]
    assert [(c.clase, c.productos, c.porcentaje) for c in r.resumen] == [("A", 2, 80), ("B", 1, 15), ("C", 1, 5)]
    assert [p.nombre for p in r.sin_ventas] == ["Épsilon"] and r.total_monto == 2400


def test_abc_umbrales_configurables(reportes):
    r = reportes.abc(ultimos_dias=30, porcentaje_a=50, porcentaje_b=30)
    assert [f.clase for f in r.productos] == ["A", "B", "C", "C"]
    with pytest.raises(ErrorValidacion):
        reportes.abc(ultimos_dias=30, porcentaje_a=90, porcentaje_b=20)


def test_abc_sin_ventas(reportes):
    r = reportes.abc(desde=date(2026, 6, 1), hasta=date(2026, 6, 30))
    assert r.sin_datos and r.productos == [] and "clasificación ABC" in r.avisos[-1]


# ------------------------------------------------- rotación y valorización

def test_valorizacion(reportes):
    r = reportes.rotacion(ultimos_dias=30)
    (total,) = r.valorizacion_total
    assert (total.moneda, total.productos, total.valor_costo, total.valor_venta, total.margen_potencial) == (
        "BOB", 5, 968, 1806, 838)
    por_grupo = {f.grupo: (f.valor_costo, f.valor_venta) for f in r.valorizacion}
    assert por_grupo == {"Herramientas": (700, 1400), "Pinturas": (108, 166), "Plomería": (160, 240)}


def test_rotacion_calculada_a_mano(reportes):
    r = reportes.rotacion(ultimos_dias=30)
    grupos = {f.grupo: f for f in r.rotacion}
    herr = grupos["Herramientas"]
    # Alfa: inicial 20 − (50 − 60) = 30 → 300; Beta: 100 − (−80) = 180 → 900. Finales 200 y 500.
    assert (herr.inventario_inicial, herr.inventario_final, herr.inventario_promedio) == (1200, 700, 950)
    assert herr.costo_vendido == 60 * 10 + 72 * 5 == 960
    assert herr.rotacion == round(960 / 950, 3) and herr.otras_salidas == 8      # la merma de Beta
    assert grupos["Plomería"].rotacion == 0 and grupos["Plomería"].dias_inventario is None
    (total,) = r.rotacion_total
    assert (total.costo_vendido, total.inventario_inicial, total.inventario_final) == (1260, 1768, 968)
    assert total.rotacion == round(1260 / 1368, 3)
    assert total.dias_inventario == round(30 / (1260 / 1368), 1)


def test_rotacion_periodo_pasado_reconstruye_el_final(reportes):
    r = reportes.rotacion(desde=date(2026, 9, 26), hasta=date(2026, 10, 15), categoria="Herramientas")
    (total,) = r.rotacion_total
    # Alfa: final = 20 − neto posterior (−30 el 20/10) = 50 → 500; Beta final 100 → 500.
    assert (total.inventario_inicial, total.inventario_final) == (1200, 1000)
    assert total.costo_vendido == 30 * 10 + 72 * 5 == 660
    assert total.rotacion == 0.6 and total.dias_inventario == round(20 / 0.6, 1)


def test_rotacion_sin_ventas_deja_valorizacion(reportes):
    r = reportes.rotacion(desde=date(2026, 6, 1), hasta=date(2026, 6, 30))
    assert r.sin_datos and all(f.rotacion is None for f in r.rotacion_total)
    assert r.valorizacion_total[0].valor_costo == 968 and "valorización sí está disponible" in r.avisos[0]


def test_avisa_precios_o_costos_en_cero(conector_reportes):
    class SinPrecio(CsvExcelConnector):
        def listar_productos(self):
            return [p.model_copy(update={"precio_venta": 0.0}) if p.id in ("P0001", "P0003") else p
                    for p in super().listar_productos()]

    r = ServicioReportes(SinPrecio(conector_reportes.data_dir), hoy=lambda: HOY).rotacion(ultimos_dias=30)
    assert any("2 producto(s) con stock tienen precio de venta 0 («Alfa», «Gamma»)" in a for a in r.avisos)
    assert r.valorizacion_total[0].valor_venta == 1806 - 400 - 150


def test_monedas_no_se_mezclan(conector_reportes):
    class ConMonedas(CsvExcelConnector):
        def listar_productos(self):
            return [p.model_copy(update={"moneda": "USD"}) if p.id == "P0002" else p for p in super().listar_productos()]

    r = ServicioReportes(ConMonedas(conector_reportes.data_dir), hoy=lambda: HOY).rotacion(ultimos_dias=30)
    assert {f.moneda: f.valor_costo for f in r.valorizacion_total} == {"BOB": 468, "USD": 500}
    assert {(f.grupo, f.moneda) for f in r.valorizacion} >= {("Herramientas", "BOB"), ("Herramientas", "USD")}


# --------------------------------------------------- resumen y exportación

def test_resumen_ejecutivo(reportes):
    r = reportes.resumen_ejecutivo()
    texto = " ".join(r.hallazgos)
    assert "BOB 2.400" in texto and "«Alfa»" in texto and "2 productos (clase A)" in texto
    assert "2 productos están bajo el stock mínimo" in texto and "«Delta»" in texto
    assert "1 productos con stock no se movieron en 90 días" in texto and "BOB 160" in texto
    assert r.productos_bajo_stock == 2 and r.capital_inmovilizado == {"BOB": 160} and not r.sin_datos


def test_resumen_sin_ventas(conector_reportes):
    r = ServicioReportes(conector_reportes, hoy=lambda: date(2027, 3, 1)).resumen_ejecutivo()
    assert r.sin_datos and r.hallazgos[0].startswith("No hay ventas registradas")
    assert any("valorizado" in h for h in r.hallazgos)


def test_exportar_reporte(reportes, tmp_path):
    r = reportes.abc(ultimos_dias=30)
    ruta = exportar_reporte(r, tmp_path / "exports", r.generado_en)
    assert ruta.name.startswith("reporte_abc_") and ruta.suffix == ".xlsx"
    libro = load_workbook(ruta)
    assert libro.sheetnames == ["Resumen ABC", "Productos", "Sin ventas", "Cómo se calculó"]
    assert [c.value for c in libro["Productos"][1]][:2] == ["producto_id", "codigo"]
    info = [tuple(f) for f in libro["Cómo se calculó"].values]
    assert any(campo == "método" and "acumulado ANTES" in valor for campo, valor in info)


# ------------------------------------------------- puerto: implementación base

def test_resumen_movimientos_por_defecto_igual_al_optimizado(conector_reportes):
    optimizado = {r.producto_id: r for r in conector_reportes.resumen_movimientos(date(2026, 9, 26), None)}
    base = {r.producto_id: r for r in InventarioConnector.resumen_movimientos(conector_reportes, date(2026, 9, 26), None)}
    assert optimizado == base and optimizado["P0001"].neto == -10


# --------------------------------------------------------- datos sintéticos

def test_reportes_con_datos_sinteticos(tmp_path):
    import importlib.util

    spec = importlib.util.spec_from_file_location("generar_datos", Path(__file__).parents[1] / "scripts" / "generar_datos.py")
    generador = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(generador)
    generador.generar(tmp_path, dias=120, inicio=date(2026, 6, 1))
    s = ServicioReportes(CsvExcelConnector(tmp_path), hoy=lambda: date(2026, 9, 28))

    abc = s.abc()
    assert sum(c.porcentaje for c in abc.resumen) == pytest.approx(100, abs=0.05)
    a = next(c for c in abc.resumen if c.clase == "A")
    assert 80 <= a.porcentaje < 80 + max(f.porcentaje for f in abc.productos)
    rot = s.rotacion()
    assert sum(f.valor_costo for f in rot.valorizacion) == pytest.approx(rot.valorizacion_total[0].valor_costo, abs=0.05)
    assert rot.rotacion_total[0].rotacion > 0 and not any("negativo" in a for a in rot.avisos)
    resumen = s.resumen_ejecutivo()
    assert len(resumen.hallazgos) >= 5 and not resumen.sin_datos


# ---------------------------------------------------------------------- MCP

@pytest.fixture
def anyio_backend() -> str:
    return "asyncio"


@pytest.mark.anyio
async def test_tools_de_reportes(conector_reportes, reloj, auditoria, tmp_path):
    from stockmind.application.confirmaciones import GestorConfirmaciones, MemoriaConfirmacionStore
    from stockmind.application.servicio import InventarioService

    servicio = InventarioService(conector_reportes, GestorConfirmaciones(MemoriaConfirmacionStore(), reloj=reloj),
                                 auditoria, exports_dir=tmp_path / "exports", hoy=lambda: HOY, reloj=reloj)
    async with Client(crear_servidor(servicio)) as client:
        async def tool(nombre, args=None):
            return (await client.call_tool(nombre, args or {})).structured_content

        abc = await tool("reporte_abc", {"ultimos_dias": 30, "exportar": True})
        assert abc["ok"] and [c["productos"] for c in abc["resumen"]] == [2, 1, 1]
        assert Path(abc["archivo"]).exists() and abc["metodo"]
        doble = await tool("reporte_mas_vendidos", {"desde": "2026-10-01", "ultimos_dias": 7})
        assert doble["ok"] is False and doble["invalidos"][0]["campo"] == "ultimos_dias"
        bajo = await tool("reporte_bajo_stock", {"categoria": "Pinturas"})
        assert [p["nombre"] for p in bajo["productos"]] == ["Delta"]
        resumen = await tool("resumen_ejecutivo")
        assert resumen["hallazgos"] and resumen["periodo"]["dias"] == 30
        assert (await tool("reporte_sin_movimiento"))["productos"][0]["nombre"] == "Épsilon"
        assert (await tool("reporte_rotacion", {"ultimos_dias": 30}))["rotacion_total"][0]["costo_vendido"] == 1260
