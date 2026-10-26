"""Exporta cualquier reporte a .xlsx en exports/.

Cada lista de filas del reporte va en su propia hoja; la última hoja, "Cómo se
calculó", tiene el título, el período, los filtros, el método y los avisos.
"""

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel

from stockmind.domain.reportes import Reporte

_NOMBRES_HOJA = {
    "por_cantidad": "Top por cantidad", "por_monto": "Top por monto", "productos": "Productos",
    "resumen": "Resumen ABC", "sin_ventas": "Sin ventas", "valorizacion_total": "Valorización total",
    "valorizacion": "Valorización por categoría", "rotacion_total": "Rotación total",
    "rotacion": "Rotación por categoría", "top_por_monto": "Top por monto", "bajo_stock": "Bajo stock",
    "abc": "Resumen ABC",
}


def _celda(valor: Any) -> Any:
    if isinstance(valor, (dict, list)):
        return ", ".join(f"{k}: {v}" for k, v in valor.items()) if isinstance(valor, dict) else ", ".join(map(str, valor))
    if isinstance(valor, datetime):
        return valor.astimezone().replace(tzinfo=None)
    return valor


def _negrita(hoja) -> None:
    for celda in hoja[1]:
        celda.font = Font(bold=True)


def exportar_reporte(reporte: Reporte, carpeta: Path, momento: datetime) -> Path:
    libro = Workbook()
    libro.remove(libro.active)
    for campo in type(reporte).model_fields:
        valor = getattr(reporte, campo)
        if not (isinstance(valor, list) and valor and isinstance(valor[0], BaseModel)):
            continue
        hoja = libro.create_sheet(_NOMBRES_HOJA.get(campo, campo)[:31])
        columnas = list(type(valor[0]).model_fields)
        hoja.append(columnas)
        for fila in valor:
            hoja.append([_celda(getattr(fila, c)) for c in columnas])
        _negrita(hoja)

    info = libro.create_sheet("Cómo se calculó")
    info.append(["campo", "valor"])
    info.append(["reporte", reporte.titulo])
    if reporte.periodo:
        info.append(["período", f"{reporte.periodo.desde:%d/%m/%Y} a {reporte.periodo.hasta:%d/%m/%Y} "
                                f"({reporte.periodo.dias} días)"])
    info.append(["generado", _celda(reporte.generado_en)])
    for k, v in reporte.filtros.items():
        info.append([f"filtro: {k}", _celda(v)])
    ya_incluidos = {"tipo", "titulo", "filtros", "periodo", "generado_en", "metodo", "avisos", "hallazgos"}
    for campo, valor in reporte:
        if campo not in ya_incluidos and isinstance(valor, (bool, int, float, str, date, dict)):
            info.append([campo, _celda(valor)])
    for nombre, lineas in (("hallazgo", getattr(reporte, "hallazgos", [])), ("método", reporte.metodo),
                           ("aviso", reporte.avisos)):
        for linea in lineas:
            info.append([nombre, linea])
    _negrita(info)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 120

    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta = carpeta / f"reporte_{reporte.tipo}_{momento.astimezone():%Y%m%d_%H%M%S}.xlsx"
    libro.save(ruta)
    return ruta
